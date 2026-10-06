"""지원 가능한 지역·경력 조건을 확인하고 한 파일로 정리합니다."""
import re
from pathlib import Path
from datetime import datetime
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

from notion_sync import employment_type, link_key, role


def eligible(job):
    location = str(job.get('location') or '')
    if not re.search(r'서울|경기', location):
        return False
    career = re.sub(r'\s+', '', str(job.get('career') or ''))
    entry = '신입' in career or '경력무관' in career
    intern = '인턴' in str(job.get('work_type') or '') + str(job.get('title') or '')
    # 경력 필수인 인턴은 제외. 신입/경력 병행 모집은 포함.
    mandatory_experience = not entry and '경력' in career and career not in ('경력없음', '경력무관')
    return entry or (intern and not mandatory_experience)


def filter_jobs(jobs):
    result = []
    seen = set()
    for job in jobs:
        key = link_key(job.get('link') or '')
        if eligible(job) and key and key not in seen:
            seen.add(key)
            result.append(job)
    return result


def excel_safe(value):
    text = str(value or '')
    # 외부 공고 문자열이 Excel 수식으로 해석되지 않도록 저장.
    return "'" + text if text.startswith(('=', '+', '-', '@')) else text


def save_excel(jobs, filename=None):
    today = datetime.now(ZoneInfo('Asia/Seoul'))
    filename = Path(filename or 'output/채용공고_' + today.strftime('%Y%m%d_%H%M') + '.xlsx')
    filename.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = '공고 목록'
    headers = ['검토 상태', '직무', '회사명', '공고명', '지역', '경력 조건',
               '채용 형태', '마감일 원문', '학력', '공고 링크', '메모']
    ws.append(headers)
    for job in jobs:
        ws.append([excel_safe(value) for value in [
            '미검토', role(job.get('keyword')), job.get('company'), job.get('title'),
            job.get('location'), job.get('career'), employment_type(job),
            job.get('deadline'), job.get('education'), job.get('link'), '',
        ]])
        cell = ws.cell(ws.max_row, 10)
        cell.hyperlink = job.get('link')
        cell.font = Font(color='0563C1', underline='single')
    ws.freeze_panes = 'E2'
    ws.auto_filter.ref = ws.dimensions
    if jobs:
        table = Table(displayName='JobPostings', ref=ws.dimensions)
        table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
        ws.add_table(table)
    widths = [13, 22, 27, 65, 24, 19, 27, 22, 15, 24, 35]
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(1, index).column_letter].width = width
    for cell in ws[1]:
        cell.fill = PatternFill('solid', fgColor='17365D')
        cell.font = Font(color='FFFFFF', bold=True)
    ws.row_dimensions[1].height = 26
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical='top', wrap_text=True)
        ws.row_dimensions[row[0].row].height = 44
    note = wb.create_sheet('읽는 방법')
    for line in [
        ['수집 시각', today.isoformat(timespec='seconds')],
        ['공고 수', len(jobs)],
        ['지역', '서울·경기도 (검색 결과의 근무지 기준)'],
        ['경력', '신입·경력무관·경력 필수가 아닌 인턴'],
        ['보는 방법', '공고 목록의 머리글 필터로 직무·지역·채용 형태를 선택하세요.'],
        ['링크', '공고 링크 셀을 클릭하면 원문으로 이동합니다.'],
        ['지원 관리', '검토 상태와 메모 열에 직접 기록할 수 있습니다.'],
        ['마감', '연도가 생략된 날짜는 추정하지 않고 원문으로 표시합니다.'],
        ['범위', '키워드당 최대 5페이지. 전체 사이트 공고를 모두 수집하는 것은 아닙니다.'],
        ['주의', '직무는 검색어 기준의 임시 분류. 정확한 업무·자격·고용 형태는 원문 확인이 필요합니다.'],
    ]:
        note.append(line)
    note.column_dimensions['A'].width = 18
    note.column_dimensions['B'].width = 95
    wb.save(filename)
    print(f'📊 조건에 맞는 {len(jobs)}개 공고를 {filename}에 저장했습니다.')
    return str(filename)
