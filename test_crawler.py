import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock, patch

from notion_sync import NotionSync, deadline_date, employment_type
from saramin_crawler import SaraminCrawler
from job_report import eligible, save_excel
from openpyxl import load_workbook


class CrawlerTests(unittest.TestCase):
    def test_searches_requested_roles_and_saves_without_email(self):
        crawler = SaraminCrawler()
        job = {'title': '교육운영 신입', 'company': '테스트', 'link': 'https://example.com/1', 'location': '서울 강남구', 'career': '신입'}
        crawler.search_jobs = Mock(return_value=[job])
        crawler.save_to_csv = Mock(return_value='test.csv')
        crawler.send_email_notification = Mock()
        with patch('saramin_crawler.NotionSync.from_environment', return_value=None), patch('saramin_crawler.save_excel'):
            jobs = crawler.run_advanced_crawler({'sender_email': None})
        keywords = {call.kwargs['keyword'] for call in crawler.search_jobs.call_args_list}
        self.assertTrue({'교육운영', '인사', '경영지원', '일반행정'} <= keywords)
        self.assertEqual(len(jobs), 1)
        crawler.save_to_csv.assert_called_once_with(jobs)
        crawler.send_email_notification.assert_not_called()

    def test_region_and_entry_filters(self):
        for location, career, title, expected in [
            ('서울 강남구', '신입', '인사', True),
            ('경기 화성시', '경력무관', '경영지원', True),
            ('서울 중구', '신입·경력', '행정', True),
            ('서울 중구', '경력3년↑', '인사', False),
            ('부산', '신입', '인사', False),
            ('서울', '경력 없음', '체험형 인턴', True),
            ('서울', '경력2년↑', '인턴', False),
            ('경기 수원시', '경력 없음', '행정', False),
        ]:
            with self.subTest(location=location, career=career, title=title):
                self.assertEqual(eligible({'location': location, 'career': career,
                                           'title': title}), expected)

    def test_workbook_filters_hyperlinks_and_formula_safety(self):
        with tempfile.TemporaryDirectory() as directory:
            path = str(Path(directory) / 'jobs.xlsx')
            save_excel([{'title': '=1+1', 'company': '기업', 'keyword': '인사',
                         'location': '서울', 'career': '신입',
                         'link': 'https://example.com/job'}], path)
            wb = load_workbook(path)
            ws = wb['공고 목록']
            self.assertEqual(ws['D2'].data_type, 's')
            self.assertEqual(ws['J2'].hyperlink.target, 'https://example.com/job')
            self.assertEqual(ws.auto_filter.ref, 'A1:K2')
            self.assertIn('JobPostings', ws.tables)
            wb.close()

    def test_csv_contains_collected_fields(self):
        with tempfile.TemporaryDirectory() as directory:
            filename = str(Path(directory) / 'jobs.csv')
            SaraminCrawler().save_to_csv([{'title': '교육운영', 'company': '기업'}], filename)
            self.assertIn('교육운영', Path(filename).read_text(encoding='utf-8-sig'))

    def test_notion_pagination_and_same_id_dedup(self):
        client = NotionSync('dummy', 'source')
        schema = {'기업명': {'type': 'title'}, '지원 링크': {'type': 'url'}}
        client.request = Mock(side_effect=[
            {'properties': schema},
            {'results': [], 'has_more': True, 'next_cursor': 'next'},
            {'results': [{'properties': {'지원 링크': {'url':
                'https://www.saramin.co.kr/zf_user/jobs/relay/view?rec_idx=1&tracking=old'}}}], 'has_more': False},
            {'id': 'new'},
        ])
        jobs = [{'company': '기업', 'title': '행정', 'keyword': '일반행정',
                 'link': 'https://www.saramin.co.kr/zf_user/jobs/relay/view?rec_idx=1'},
                {'company': '새기업', 'title': '인사', 'keyword': '인사',
                 'link': 'https://www.saramin.co.kr/zf_user/jobs/relay/view?rec_idx=2'}]
        self.assertEqual(client.sync(jobs), {'created': 1, 'skipped': 1})
        self.assertEqual(client.request.call_args_list[2].kwargs['json']['start_cursor'], 'next')
        payload = client.request.call_args.kwargs['json']
        self.assertEqual(payload['parent']['data_source_id'], 'source')
        self.assertNotIn('상태', payload['properties'])

    def test_uncertain_dates_and_employment_are_not_invented(self):
        self.assertIsNone(deadline_date('~ 10/31(토)'))
        self.assertIsNone(deadline_date('채용시'))
        self.assertEqual(deadline_date('~2026.10.31'), '2026-10-31')
        self.assertIsNone(deadline_date('2026.02.30'))
        self.assertEqual(employment_type({'title': '인턴'}), '인턴(전환 여부 미확인)')
        self.assertEqual(employment_type({'title': '채용전환형 인턴'}), '인턴(채용전환형)')
        self.assertEqual(employment_type({'work_type': '정규직·계약직'}), '확인 필요')


if __name__ == '__main__':
    unittest.main()
